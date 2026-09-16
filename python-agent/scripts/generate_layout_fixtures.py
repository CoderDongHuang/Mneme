from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


FIXTURES = Path(__file__).resolve().parents[1] / "evaluation" / "fixtures"


def scanned_page() -> Path:
    image = Image.new("RGB", (1400, 1900), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=36)
    draw.text(
        (100, 180),
        "SCAN-ALPHA: handwritten archive retention is 42 months.\n"
        "OCR confidence and page coordinates must remain attached.",
        fill="black",
        font=font,
        spacing=24,
    )
    output = FIXTURES / "scanned-page.png"
    image.save(output, format="PNG", optimize=False)
    return output


def build_pdf() -> None:
    width, height = A4
    document = canvas.Canvas(str(FIXTURES / "complex-layout.pdf"), pagesize=A4, invariant=1)
    document.setTitle("Mneme annotated layout regression fixture")
    document.setAuthor("Mneme")
    scan = scanned_page()
    document.drawImage(ImageReader(str(scan)), 0, 0, width=width, height=height)
    document.showPage()

    document.setFont("Helvetica", 10)
    for line in range(7):
        document.drawString(42, height - 65 - line * 45, "COLUMN-LEFT")
        document.drawString(42, height - 80 - line * 45, "The left policy sets the recovery point objective to 15 minutes.")
        document.drawString(316, height - 65 - line * 45, "COLUMN-RIGHT")
        document.drawString(316, height - 80 - line * 45, "The right policy sets the recovery time objective to 60 minutes.")
    document.showPage()

    for part, rows in ((1, range(1, 9)), (2, range(9, 17))):
        document.setFont("Helvetica-Bold", 13)
        document.drawString(42, height - 45, f"CROSS-PAGE-TABLE part {part}")
        y = height - 80
        document.setFont("Helvetica", 10)
        for row in rows:
            document.rect(42, y - 24, 511, 32)
            document.drawString(50, y - 12, f"Item {row:02d} | value={row * 7} | owner=team-{row % 3}")
            y -= 32
        if part == 2:
            document.drawString(42, 190, "FORMULA: quality = (supported claims / total claims) * 100")
            document.setFillColorRGB(0.35, 0.45, 0.25)
            for index, bar_height in enumerate((45, 90, 135, 180)):
                document.rect(330 + index * 45, 45, 30, bar_height, fill=1)
            document.setFillColorRGB(0, 0, 0)
        document.showPage()
    document.save()


def build_workbook() -> None:
    workbook = Workbook()
    workbook.properties.created = datetime(2026, 9, 16, tzinfo=timezone.utc)
    sheet = workbook.active
    sheet.title = "FormulaChart"
    sheet.append(["Month", "Reviewed", "Correct", "Accuracy"])
    for month, reviewed, correct in (("Jan", 20, 14), ("Feb", 25, 20), ("Mar", 30, 27)):
        row = sheet.max_row + 1
        sheet.append([month, reviewed, correct, f"=C{row}/B{row}"])
    chart = BarChart()
    chart.title = "Reviewed vs Correct"
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    sheet.add_chart(chart, "F2")
    workbook.save(FIXTURES / "formula-chart.xlsx")


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_workbook()


if __name__ == "__main__":
    main()
